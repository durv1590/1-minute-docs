"""Builds a DOCX from OCR'd page texts. Usage: text_to_docx.py PAGES.json OUTPUT.docx

PAGES.json: {"pages": ["page 1 text", ...]}. Blank lines separate paragraphs; pages are separated by page breaks.
Devanagari text gets a complex-script font hint so Word picks a Hindi-capable font.
"""
import json
import sys

from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml.ns import qn
from docx.shared import Pt


def main() -> int:
    src, dst = sys.argv[1], sys.argv[2]
    with open(src, encoding="utf-8") as fh:
        pages = json.load(fh)["pages"]
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    rfonts = style.element.get_or_add_rPr().get_or_add_rFonts()
    rfonts.set(qn("w:cs"), "Nirmala UI")
    for index, text in enumerate(pages):
        paragraphs = [p.strip() for p in text.replace("\r", "").split("\n\n")]
        for para in paragraphs:
            if para:
                doc.add_paragraph(" ".join(line.strip() for line in para.split("\n")))
        if index < len(pages) - 1:
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    doc.save(dst)
    print(json.dumps({"ok": True}))
    return 0


if __name__ == "__main__":
    try:
        sys.exit(main())
    except Exception as exc:  # noqa: BLE001
        print(json.dumps({"error": type(exc).__name__}))
        sys.exit(1)
