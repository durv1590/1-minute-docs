"""Rebuilds an editable DOCX from a text PDF's text layer. Usage: pdf_text_to_docx.py BBOX.xhtml OUTPUT.docx

BBOX.xhtml is the output of `pdftotext -bbox-layout` (poppler). Each text block becomes a paragraph; blocks whose line
height is clearly larger than the body text become headings; pages are separated by page breaks. Poppler extracts complex
scripts such as Devanagari in logical order, which LibreOffice's PDF import does not. Tables, images, columns and exact
positions are not reconstructed (reported as a warning by the caller).
Uses python-docx (MIT); no AGPL component. Prints one JSON line: {"ok": true, "paragraphs": N, "headings": M}.
"""
import json
import statistics
import sys
import xml.etree.ElementTree as ET

from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml.ns import qn
from docx.shared import Pt

NS = {"x": "http://www.w3.org/1999/xhtml"}


def line_text(line: ET.Element) -> str:
    return " ".join((w.text or "").strip() for w in line.findall("x:word", NS) if (w.text or "").strip())


def line_height(line: ET.Element) -> float:
    return float(line.get("yMax", 0)) - float(line.get("yMin", 0))


def read_pages(path: str) -> list:
    root = ET.parse(path).getroot()
    pages = []
    for page in root.iter(f"{{{NS['x']}}}page"):
        blocks = []
        for block in page.iter(f"{{{NS['x']}}}block"):
            lines = [ln for ln in block.findall("x:line", NS) if line_text(ln)]
            if not lines:
                continue
            text = ""
            for ln in lines:
                t = line_text(ln)
                # Re-join words hyphenated across lines ("docu-" + "ment"); other lines are joined with a space.
                if text.endswith("-") and t[:1].islower():
                    text = text[:-1] + t
                else:
                    text = f"{text} {t}" if text else t
            blocks.append({"text": text, "height": statistics.median(line_height(ln) for ln in lines), "lines": len(lines)})
        pages.append(blocks)
    return pages


def main() -> int:
    src, dst = sys.argv[1], sys.argv[2]
    pages = read_pages(src)
    heights = [b["height"] for blocks in pages for b in blocks for _ in range(b["lines"])]
    body = statistics.median(heights) if heights else 0
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    style.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:cs"), "Nirmala UI")
    paragraphs = headings = 0
    for index, blocks in enumerate(pages):
        for block in blocks:
            ratio = block["height"] / body if body else 1
            if block["lines"] <= 2 and ratio >= 1.35:
                doc.add_heading(block["text"], level=1 if ratio >= 1.8 else 2)
                headings += 1
            else:
                doc.add_paragraph(block["text"])
                paragraphs += 1
        if index < len(pages) - 1:
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    doc.save(dst)
    print(json.dumps({"ok": True, "paragraphs": paragraphs, "headings": headings}))
    return 0


if __name__ == "__main__":
    try:
        sys.exit(main())
    except Exception as exc:  # noqa: BLE001
        print(json.dumps({"ok": False, "error": type(exc).__name__}))
        sys.exit(1)
